"""Build PowerPoint presentations from parsed song data.

Slide structure (all coordinates in inches):
  ┌──────────────────────────────────────────────────────────┐
  │  [Title Chinese]                   [Title Pinyin]         │
  ├──────────────────────────────────────────────────────────┤
  │  Col 0                   Col 1                            │
  │  [#] [section 1 image]   [#] [section 2 image]           │
  │                              [section 3 image]            │
  └──────────────────────────────────────────────────────────┘

  • Each column is packed independently (greedy top-to-bottom).
  • All section images on a slide use the same scale so text is uniform.
  • Verse numbers are native pptx text boxes to the left of section images.
  • Chorus sections have a '* Chorus 副歌:' / '* Chorus:' label baked into
    the top of their PIL image; other section types get no label.
  • Non-chorus sections followed by a chorus in source order have trailing
    asterisk(s) baked into the bottom of their PIL image.
  • Duplicate chorus sections (same content) are deduplicated to appear once
    at their first occurrence position; later duplicates are skipped.
  • Section images have transparent backgrounds.
  • English sections are placed as native pptx text boxes (no PIL images).
"""

from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from .config import SlideConfig, Song
from .renderer import render_section, char_height_px

# ── Slide geometry (inches) ───────────────────────────────────────────────────
_SLIDE_W = 13.333
_SLIDE_H = 7.5

_MARGIN_L = 0.50
_MARGIN_R = 0.50
_MARGIN_TOP = 0.20

_IMG_NO_NUM_INSET = 0.10

_TITLE_H   = 0.65
_TITLE_GAP = 0.15

# Reserve space in the top-right for future book/page metadata.
_TOPRIGHT_RESERVE = 2.0

_CONTENT_TOP = _MARGIN_TOP + _TITLE_H + _TITLE_GAP
_CONTENT_H   = _SLIDE_H - _CONTENT_TOP - 0.30
_CONTENT_W   = _SLIDE_W - _MARGIN_L - _MARGIN_R

_COL_GAP     = 0.20
_ROW_GAP     = 0.30
_VERSE_NUM_W = 0.45

# Minimum character height on Chinese slides (inches). Derived from a 28pt
# floor: 28 / 72 ≈ 0.389", but the PIL render uses a tight bbox without
# leading, so we scale down accordingly. This matches the 28pt English floor.
_MIN_CHAR_PT  = 28
_MIN_CHAR_H_IN = _MIN_CHAR_PT / 72 * 0.7   # ≈ 0.272"

# Title / verse-number font sizes (pt)
_TITLE_ZH_PT  = 36
_TITLE_PY_PT  = 36
_VERSE_NUM_PT = 22

# ── English text constants ────────────────────────────────────────────────────
_MIN_ENGLISH_PT = 28
_MAX_ENGLISH_PT = 40
_ENGLISH_LINE_H  = 1.35
_ENGLISH_CHAR_W  = 0.52


# ── Chorus / section helpers ──────────────────────────────────────────────────

def _chorus_label(language: str) -> str:
    return '* Chorus 副歌:' if language == 'chinese' else '* Chorus:'


def _section_content_key(section) -> tuple:
    """Hashable key for section line content — used for chorus deduplication."""
    return tuple(
        tuple(tok.char for word in line.words for tok in word.tokens)
        for line in section.lines
    )


def _split_single_section_for_columns(display_sections, labels, trailing_asts,
                                        n_cols: int) -> tuple:
    """Split a lone section's lines across `n_cols` pseudo-sections so a long
    single verse can be rendered across multiple columns at a larger scale."""
    from .config import Section
    section = display_sections[0]
    lines = section.lines
    n = len(lines)
    if n < n_cols:
        return display_sections, labels, trailing_asts

    # Roughly-even split, with earlier chunks getting extra lines when uneven.
    chunk_sizes = [n // n_cols] * n_cols
    for i in range(n % n_cols):
        chunk_sizes[i] += 1

    pieces = []
    start = 0
    for i, size in enumerate(chunk_sizes):
        end = start + size
        number = section.number if i == 0 else None
        pieces.append(Section(lines=lines[start:end],
                              type=section.type, number=number))
        start = end

    new_labels = [labels[0]] + [''] * (len(pieces) - 1)
    new_asts = [trailing_asts[0]] + [0] * (len(pieces) - 1)
    return pieces, new_labels, new_asts


def _verse_and_chorus_indices(display_sections, chorus_ref_map):
    """Split display indices into verses (in source order) and the unique
    chorus indices each verse references.

    Returns:
        verse_d_idxs: list[int] of display indices that are verses, in source
                      order. Verses that DON'T reference a chorus are still
                      included.
        chorus_for_verse: list[int|None], same length as verse_d_idxs; the
                          referenced chorus display index, or None.
    """
    verse_d_idxs = []
    chorus_for_verse = []
    referenced_choruses = set()
    for refs in chorus_ref_map.values():
        referenced_choruses.update(refs)

    for d_idx in range(len(display_sections)):
        if d_idx in referenced_choruses:
            # Skip standalone chorus entries — they get placed per-slide.
            continue
        verse_d_idxs.append(d_idx)
        refs = chorus_ref_map.get(d_idx, ())
        chorus_for_verse.append(refs[0] if refs else None)
    return verse_d_idxs, chorus_for_verse


def _place_chorus_on_slide(verse_cols, chorus_idx, fits_predicate,
                            n_cols: int) -> tuple:
    """Insert `chorus_idx` into `verse_cols` per the user-specified rules:

      1. Append to col 0 (chorus follows verses in source order on left).
      2. If that doesn't fit: chorus is the FIRST section in col 1; verses
         from col 1 are pulled back into col 0 if they fit there.
      3. If chorus + col-1 verses still don't fit: chorus goes at col 1 bottom.

    `fits_predicate(cols) -> bool` is the language-specific fit check.
    Returns (cols_with_chorus, success).
    """
    col0 = list(verse_cols[0]) if verse_cols else []
    col1 = list(verse_cols[1]) if len(verse_cols) > 1 else []

    if n_cols == 1:
        cand = [col0 + [chorus_idx]]
        return (cand, fits_predicate(cand))

    # Strategy 1: append chorus to col 0 (after the verses there).
    cand = [col0 + [chorus_idx], col1]
    if fits_predicate(cand):
        return (cand, True)

    # Strategy 2: chorus is the first section in col 1, with verses
    # rebalanced — try pulling all col 1 verses back into col 0 first.
    all_verses = col0 + col1
    for k in range(len(all_verses), -1, -1):
        # k verses in col 0, chorus + remainder in col 1
        cand = [all_verses[:k], [chorus_idx] + all_verses[k:]]
        if fits_predicate(cand):
            return (cand, True)

    # Strategy 3: chorus at col 1 bottom (after existing col 1 verses).
    cand = [col0, col1 + [chorus_idx]]
    if fits_predicate(cand):
        return (cand, True)

    return (verse_cols, False)


def _pack_chinese_with_chorus(verse_d_idxs, chorus_for_verse,
                                section_imgs, config: SlideConfig,
                                img_w: float, ch_px: int,
                                content_h: float, max_rows,
                                min_char_h: float = None) -> list:
    """Pack verses into slides; for each slide that contains verses
    referencing a chorus, place the chorus per `_place_chorus_on_slide`.

    Each slide shows the chorus AT MOST ONCE, regardless of how many of its
    referencing verses are on that slide.
    """
    if min_char_h is None:
        min_char_h = _MIN_CHAR_H_IN
    n_cols = config.columns
    n = len(verse_d_idxs)

    def fits(cols):
        non_empty = [c for c in cols if c]
        if not non_empty:
            return True
        if max_rows and any(len(c) > max_rows for c in non_empty):
            return False
        # Use the same uniform scale the renderer will apply so the fit
        # check matches what is actually rendered. Pass only non-empty columns
        # to avoid division-by-zero in _uniform_scale_for_slide.
        scale = _uniform_scale_for_slide(non_empty, section_imgs, img_w, content_h)
        return ch_px * scale >= min_char_h

    slides = []
    i = 0
    while i < n:
        chorus_idx = chorus_for_verse[i]

        # Verses without a chorus reference: pack normally and continue.
        if chorus_idx is None:
            # Find run of consecutive verses with no chorus reference.
            j = i
            while j < n and chorus_for_verse[j] is None:
                j += 1
            sub_imgs = [section_imgs[d] for d in verse_d_idxs[i:j]]
            sub_slides = _pack_into_columns(sub_imgs, config, img_w, ch_px,
                                              content_h, min_char_h)
            for sub in sub_slides:
                slides.append([[verse_d_idxs[i + k] for k in col] for col in sub])
            i = j
            continue

        # Find the upper bound: how many verses fit in n_cols columns
        # (without chorus) starting at i?
        sub_imgs_all = [section_imgs[d] for d in verse_d_idxs[i:]]
        upper = _best_pack_slide(sub_imgs_all, config, img_w, ch_px,
                                  0, len(sub_imgs_all), n_cols, max_rows,
                                  content_h, min_char_h)
        upper_count = sum(len(c) for c in upper)
        upper_count = max(1, upper_count)

        # Try N from upper_count down to 1; pick the largest N where the
        # chorus can be placed somewhere.
        chosen = None
        for N in range(upper_count, 0, -1):
            verse_subset = verse_d_idxs[i:i + N]
            sub_imgs = [section_imgs[d] for d in verse_subset]
            verse_cols_local = _best_pack_slide(sub_imgs, config, img_w, ch_px,
                                                  0, N, n_cols, max_rows,
                                                  content_h, min_char_h)
            if sum(len(c) for c in verse_cols_local) < N:
                continue  # Packer couldn't fit all N verses — try smaller.
            verse_cols_disp = [[verse_subset[j] for j in col]
                                for col in verse_cols_local]
            cand, ok = _place_chorus_on_slide(verse_cols_disp, chorus_idx,
                                                fits, n_cols)
            if ok:
                chosen = (cand, N)
                break

        if chosen is None:
            # Force one verse + chorus on this slide even if it overflows.
            forced = [[verse_d_idxs[i]], [chorus_idx]] if n_cols >= 2 \
                      else [[verse_d_idxs[i], chorus_idx]]
            chosen = (forced, 1)

        slides.append(chosen[0])
        i += chosen[1]

    return slides


def _pack_english_with_chorus(verse_d_idxs, chorus_for_verse,
                                display_sections, labels, trailing_asts,
                                n_cols: int, max_rows,
                                text_width_in: float, content_h: float) -> list:
    """English analogue of `_pack_chinese_with_chorus`. Uses visual-line
    counting for the fit predicate."""

    def fits(cols):
        # Pick the largest pt that fits all columns simultaneously, then
        # check it's at or above the floor.
        for col in cols:
            if not col:
                continue
            if max_rows and len(col) > max_rows:
                return False
            pt = _compute_english_col_font_pt(col, display_sections, labels,
                                                trailing_asts, text_width_in,
                                                content_h)
            line_h_in = pt * _ENGLISH_LINE_H / 72.0
            vis = _english_col_visual_lines(col, display_sections, labels,
                                              trailing_asts, pt, text_width_in)
            if pt <= _MIN_ENGLISH_PT and vis * line_h_in > content_h:
                return False
        # Also require the *uniform* font for the slide (min across cols)
        # to leave each col's content within bounds — already handled above
        # since each col's pt is computed independently and we check the
        # min via rendering. Conservative: re-check at uniform pt.
        non_empty = [c for c in cols if c]
        if non_empty:
            uniform_pt = min(
                _compute_english_col_font_pt(c, display_sections, labels,
                                               trailing_asts, text_width_in,
                                               content_h)
                for c in non_empty
            )
            line_h_in = uniform_pt * _ENGLISH_LINE_H / 72.0
            for c in non_empty:
                vis = _english_col_visual_lines(c, display_sections, labels,
                                                  trailing_asts, uniform_pt,
                                                  text_width_in)
                if vis * line_h_in > content_h + 1e-6:
                    return False
        return True

    n = len(verse_d_idxs)
    slides = []
    i = 0
    asts_zero = [0] * len(display_sections)  # for verse-only sub-pack
    while i < n:
        chorus_idx = chorus_for_verse[i]

        if chorus_idx is None:
            # Pack consecutive no-chorus verses normally.
            j = i
            while j < n and chorus_for_verse[j] is None:
                j += 1
            sub_d = verse_d_idxs[i:j]
            sub_slides = _pack_english_into_slides(
                [display_sections[d] for d in sub_d],
                [labels[d] for d in sub_d],
                n_cols, max_rows, text_width_in, content_h)
            for sub in sub_slides:
                slides.append([[sub_d[k] for k in col] for col in sub])
            i = j
            continue

        # Upper bound: max verses (in source order from i) that fit without
        # chorus, packed into n_cols columns.
        sub_d_all = verse_d_idxs[i:]
        sub_pack_all = _pack_english_slide(
            [display_sections[d] for d in sub_d_all],
            [labels[d] for d in sub_d_all],
            asts_zero,
            0, len(sub_d_all), n_cols, max_rows,
            text_width_in, content_h)
        upper_count = max(1, sum(len(c) for c in sub_pack_all))

        chosen = None
        for N in range(upper_count, 0, -1):
            verse_subset = verse_d_idxs[i:i + N]
            sub_pack = _pack_english_slide(
                [display_sections[d] for d in verse_subset],
                [labels[d] for d in verse_subset],
                asts_zero,
                0, N, n_cols, max_rows, text_width_in, content_h)
            if sum(len(c) for c in sub_pack) < N:
                continue
            verse_cols_disp = [[verse_subset[j] for j in col]
                                for col in sub_pack]
            cand, ok = _place_chorus_on_slide(verse_cols_disp, chorus_idx,
                                                fits, n_cols)
            if ok:
                chosen = (cand, N)
                break

        if chosen is None:
            forced = [[verse_d_idxs[i]], [chorus_idx]] if n_cols >= 2 \
                      else [[verse_d_idxs[i], chorus_idx]]
            chosen = (forced, 1)

        slides.append(chosen[0])
        i += chosen[1]

    return slides


def _compute_display_info(song: Song, dedup_chorus: bool) -> tuple:
    """Build the display section list with chorus handling and rendering metadata.

    In repeat mode (``dedup_chorus=False``), every section is emitted verbatim
    with no label and no trailing asterisks.

    In dedup mode, duplicate chorus sections (same content as an earlier
    occurrence) are removed; the first occurrence gets a ``* Chorus:`` label,
    and a verse that was followed by one or more chorus sections in the source
    gets that many trailing ``*`` markers.

    Returns:
        display_sections  – list of Section objects
        labels            – list[str]: chorus label or '' per display section
        trailing_asts     – list[int]: trailing '*' count per display section
        chorus_ref_map    – dict[verse_display_idx -> [chorus_display_idx, ...]]
                            mapping each verse with trailing asterisks to the
                            chorus body it references. Empty in repeat mode.
    """
    src = song.sections
    n = len(src)

    if not dedup_chorus:
        return (list(src), [''] * n, [0] * n, {})

    # For each non-chorus section, collect the consecutive chorus-type sections
    # that immediately follow it in source order.
    following_choruses = [[] for _ in range(n)]
    for i in range(n):
        if src[i].type not in ('chorus', 'refrain'):
            j = i + 1
            while j < n and src[j].type in ('chorus', 'refrain'):
                following_choruses[i].append(j)
                j += 1

    chorus_lbl = _chorus_label(song.language)
    seen_chorus_keys: dict = {}  # key -> display_idx of first occurrence
    src_to_display: dict = {}    # source idx -> display idx (choruses mapped to first occurrence)
    display_sections = []
    labels = []
    trailing_asts = []

    for i, section in enumerate(src):
        if section.type in ('chorus', 'refrain'):
            key = _section_content_key(section)
            if key in seen_chorus_keys:
                # duplicate — point this source idx at the first-occurrence display idx
                src_to_display[i] = seen_chorus_keys[key]
                continue
            d_idx = len(display_sections)
            seen_chorus_keys[key] = d_idx
            src_to_display[i] = d_idx
            display_sections.append(section)
            labels.append(chorus_lbl)
            trailing_asts.append(0)
        else:
            d_idx = len(display_sections)
            src_to_display[i] = d_idx
            display_sections.append(section)
            labels.append('')
            trailing_asts.append(len(following_choruses[i]))

    chorus_ref_map: dict = {}
    for i, section in enumerate(src):
        if section.type in ('chorus', 'refrain'):
            continue
        refs = [src_to_display[j] for j in following_choruses[i]]
        if refs:
            chorus_ref_map[src_to_display[i]] = refs

    return display_sections, labels, trailing_asts, chorus_ref_map


# ── Public API ────────────────────────────────────────────────────────────────

def build_presentation(song: Song, config: SlideConfig) -> bytes:
    """Build a .pptx file from a parsed Song."""
    prs = Presentation()
    prs.slide_width  = Inches(_SLIDE_W)
    prs.slide_height = Inches(_SLIDE_H)
    blank_layout = prs.slide_layouts[6]

    if song.language == 'english':
        _add_english_slides(prs, blank_layout, song, config)
    else:
        _add_chinese_slides(prs, blank_layout, song, config)

    output = BytesIO()
    prs.save(output)
    return output.getvalue()


def build_deck(song_configs: list, global_config: SlideConfig) -> bytes:
    """Build a .pptx from multiple songs."""
    prs = Presentation()
    prs.slide_width  = Inches(_SLIDE_W)
    prs.slide_height = Inches(_SLIDE_H)
    blank_layout = prs.slide_layouts[6]

    for song, overrides in song_configs:
        config = apply_config_overrides(global_config, overrides)
        if song.language == 'english':
            _add_english_slides(prs, blank_layout, song, config)
        else:
            _add_chinese_slides(prs, blank_layout, song, config)

    output = BytesIO()
    prs.save(output)
    return output.getvalue()


def apply_config_overrides(base_config: SlideConfig, overrides: dict) -> SlideConfig:
    """Apply per-song config overrides (string values) to a SlideConfig copy."""
    import copy
    cfg = copy.copy(base_config)
    def _to_bool(v):
        return str(v).strip().lower() in ('true', '1', 'yes', 'on')

    mapping = {
        'columns':          ('columns',              int),
        'rows':             ('rows_per_column',       int),
        'pinyin_size':      ('pinyin_font_size',      int),
        'char_size':        ('char_font_size',        int),
        'text_color':       ('text_color',            str),
        'max_lines':        ('max_lines_per_verse',   int),
        'english_font_size':('english_font_size_pt',  int),
        'dedup_chorus':     ('dedup_chorus',          _to_bool),
        'book':             ('book',                   str),
        'page':             ('page',                   str),
        'min_char_pt':      ('min_char_pt',            int),
    }
    for key, value in overrides.items():
        if key == 'language':
            continue
        if key in mapping:
            attr, typ = mapping[key]
            try:
                setattr(cfg, attr, typ(value))
            except (ValueError, TypeError):
                pass
    return cfg


# ── Chinese slides ────────────────────────────────────────────────────────────

def _add_chinese_slides(prs, blank_layout, song: Song, config: SlideConfig):
    """Add Chinese pinyin slides to an existing Presentation object."""
    display_sections, labels, trailing_asts, chorus_ref_map = _compute_display_info(
        song, config.dedup_chorus)

    # If the song is a single section and we have multiple columns, split the
    # section across columns so the text can be larger.
    if (config.columns >= 2
            and len(display_sections) == 1
            and len(display_sections[0].lines) >= config.columns * 4
            and not chorus_ref_map):
        display_sections, labels, trailing_asts = _split_single_section_for_columns(
            display_sections, labels, trailing_asts, config.columns)
    section_imgs = [
        render_section(s, config, section_label=labels[i],
                       trailing_asterisks=trailing_asts[i])
        for i, s in enumerate(display_sections)
    ]

    cols  = config.columns
    col_w = (_CONTENT_W - _COL_GAP * (cols - 1)) / cols
    img_w_conservative = col_w - _VERSE_NUM_W
    ch_px = char_height_px(config)

    # Per-slide content top can shift down when the title wraps. Pass the
    # available height to packing so we don't pack more than fits.
    content_top = _content_top_for_title(song.title_zh, song.title_py,
                                          config.book, config.page)
    content_h   = _SLIDE_H - content_top - 0.30

    # Per-song minimum character height override (inches).
    effective_min_char_h = (
        (config.min_char_pt / 72 * 0.7) if config.min_char_pt else _MIN_CHAR_H_IN
    )

    if chorus_ref_map:
        # Dedup mode: pack verses, place chorus per slide as a floater so
        # each slide shows the chorus once at the position dictated by the
        # user's rules (col 0 bottom → col 1 top → col 1 bottom).
        verse_d_idxs, chorus_for_verse = _verse_and_chorus_indices(
            display_sections, chorus_ref_map)
        slides = _pack_chinese_with_chorus(
            verse_d_idxs, chorus_for_verse, section_imgs, config,
            img_w_conservative, ch_px, content_h, config.rows_per_column,
            effective_min_char_h)
    else:
        slides = _pack_into_columns(section_imgs, config,
                                     img_w_conservative, ch_px,
                                     content_h, effective_min_char_h)

    for slide_cols in slides:
        slide = prs.slides.add_slide(blank_layout)

        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        if song.title_zh or song.title_py:
            _add_title_box(slide, song.title_zh, song.title_py,
                           config.text_color, config.book, config.page)
        _add_book_page_box(slide, config.book, config.page, config.text_color)

        all_section_idxs = [si for col in slide_cols for si in col]
        any_numbered = any(display_sections[si].number for si in all_section_idxs)
        num_margin   = _VERSE_NUM_W if any_numbered else 0.0
        img_inset    = 0.0 if any_numbered else _IMG_NO_NUM_INSET
        img_w        = col_w - num_margin - img_inset

        scale = _uniform_scale_for_slide(slide_cols, section_imgs, img_w, content_h)

        for col_idx, col_group in enumerate(slide_cols):
            col_left = _MARGIN_L + col_idx * (col_w + _COL_GAP)
            row_top  = content_top

            for section_idx in col_group:
                section = display_sections[section_idx]
                s_img   = section_imgs[section_idx]
                img_h   = s_img.height * scale

                if section.number:
                    _add_text_box(slide, f"{section.number}.",
                                  left=col_left, top=row_top,
                                  width=_VERSE_NUM_W, height=img_h,
                                  font_pt=_VERSE_NUM_PT, align=PP_ALIGN.LEFT,
                                  color=config.text_color, v_anchor='top')

                _place_verse_image(slide, s_img,
                                   left=col_left + num_margin + img_inset,
                                   top=row_top,
                                   scale=scale)

                row_top += img_h + _ROW_GAP


# ── English slides ────────────────────────────────────────────────────────────

def _add_english_slides(prs, blank_layout, song: Song, config: SlideConfig):
    """Add English-language slides to the presentation."""
    display_sections, labels, trailing_asts, chorus_ref_map = _compute_display_info(
        song, config.dedup_chorus)

    if not display_sections:
        return

    n_cols = _auto_detect_english_cols(display_sections, config)
    max_rows = config.rows_per_column

    # Per-slide content top can shift down when the title wraps. Compute once
    # for this song; all its slides share the same title.
    content_top = _content_top_for_title(song.title_zh, '',
                                          config.book, config.page)
    content_h   = _SLIDE_H - content_top - 0.30

    # Text width depends on column count; compute once for this song.
    col_w = (_CONTENT_W - _COL_GAP * (n_cols - 1)) / n_cols
    any_numbered_any = any(s.number for s in display_sections)
    num_margin_default = _VERSE_NUM_W if any_numbered_any else 0.0
    text_width = col_w - num_margin_default

    if chorus_ref_map:
        verse_d_idxs, chorus_for_verse = _verse_and_chorus_indices(
            display_sections, chorus_ref_map)
        slides = _pack_english_with_chorus(
            verse_d_idxs, chorus_for_verse,
            display_sections, labels, trailing_asts,
            n_cols, max_rows, text_width, content_h)
    else:
        slides = _pack_english_into_slides(display_sections, labels,
                                            n_cols, max_rows,
                                            text_width, content_h)

    for slide_cols in slides:
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        if song.title_zh:
            _add_title_box(slide, song.title_zh, '', config.text_color,
                           config.book, config.page)
        _add_book_page_box(slide, config.book, config.page, config.text_color)

        n_cols_slide = len(slide_cols)
        slide_col_w = (_CONTENT_W - _COL_GAP * (n_cols_slide - 1)) / n_cols_slide

        any_numbered = any(
            display_sections[si].number
            for col in slide_cols for si in col
        )
        num_margin = _VERSE_NUM_W if any_numbered else 0.0
        slide_text_width = slide_col_w - num_margin

        # Pick a single uniform font size that fits the worst column on this
        # slide, accounting for visual line wrap.
        font_pt = min(
            _compute_english_col_font_pt(col, display_sections, labels,
                                          trailing_asts, slide_text_width,
                                          content_h)
            for col in slide_cols if col
        )

        for col_idx, col_group in enumerate(slide_cols):
            if not col_group:
                continue
            col_left = _MARGIN_L + col_idx * (slide_col_w + _COL_GAP)

            _render_english_column(
                slide, display_sections, col_group,
                labels=labels, trailing_asts=trailing_asts,
                col_left=col_left, num_margin=num_margin,
                text_width=slide_text_width,
                font_pt=font_pt, config=config,
                content_top=content_top,
                content_h=content_h,
            )


def _auto_detect_english_cols(sections, config):
    """Return appropriate column count for English lyrics."""
    if config.columns <= 1:
        return 1

    max_chars = max(
        (len(token.char)
         for section in sections
         for line in section.lines
         for word in line.words
         for token in word.tokens
         if token.char),
        default=0,
    )
    if max_chars == 0:
        return 1

    two_col_w = (_CONTENT_W - _COL_GAP) / 2
    line_w_in = max_chars * _MIN_ENGLISH_PT * _ENGLISH_CHAR_W / 72

    if line_w_in <= two_col_w:
        return min(2, config.columns)
    return 1


def _english_section_line_count(section, label: str) -> int:
    """Total source-line count for one English section (label + content)."""
    return len(section.lines) + (1 if label else 0)


def _english_line_text(section, line_idx: int, n_asts: int) -> str:
    """Reconstruct the rendered text of one source line, including any
    trailing asterisks appended to the last line."""
    line = section.lines[line_idx]
    text = line.words[0].tokens[0].char if line.words else ''
    if line_idx == len(section.lines) - 1 and n_asts > 0:
        text = text + ' ' + ' '.join(['*'] * n_asts)
    return text


def _english_visual_lines_for_line(text: str, font_pt: float,
                                     text_width_in: float) -> int:
    """Estimate how many visual lines a single source line will wrap into.

    Uses the existing `_ENGLISH_CHAR_W` heuristic (Arial-ish 0.52em). Returns
    at least 1, even for an empty line.
    """
    import math
    if text_width_in <= 0:
        return 1
    char_w_in = font_pt * _ENGLISH_CHAR_W / 72.0
    line_w_in = max(1, len(text)) * char_w_in
    return max(1, int(math.ceil(line_w_in / text_width_in)))


def _english_section_visual_lines(section, label: str, n_asts: int,
                                    font_pt: float, text_width_in: float) -> int:
    """Total visual (post-wrap) line count for one English section at the
    given font and column width."""
    total = 1 if label else 0
    for i in range(len(section.lines)):
        text = _english_line_text(section, i, n_asts)
        total += _english_visual_lines_for_line(text, font_pt, text_width_in)
    return total


def _english_col_visual_lines(col_group, sections, labels, trailing_asts,
                                font_pt: float, text_width_in: float) -> int:
    """Total visual lines for a column (sum of per-section visual lines plus
    one blank gap between sections)."""
    n_sections = len(col_group)
    if n_sections == 0:
        return 0
    total = sum(
        _english_section_visual_lines(sections[vi], labels[vi],
                                       trailing_asts[vi],
                                       font_pt, text_width_in)
        for vi in col_group
    )
    total += max(0, n_sections - 1)  # one-line gap between sections
    return total


def _compute_english_col_font_pt(col_group, sections, labels, trailing_asts,
                                   text_width_in: float, content_h: float) -> float:
    """Largest font size (within [_MIN_ENGLISH_PT, _MAX_ENGLISH_PT]) at which
    the column's content — accounting for line wrap — fits inside `content_h`.
    """
    if not col_group:
        return _MIN_ENGLISH_PT
    # Walk pt downward in 1pt steps from MAX to MIN; pick the largest that fits.
    for pt in range(_MAX_ENGLISH_PT, _MIN_ENGLISH_PT - 1, -1):
        line_h_in = pt * _ENGLISH_LINE_H / 72.0
        vis = _english_col_visual_lines(col_group, sections, labels,
                                         trailing_asts, pt, text_width_in)
        if vis * line_h_in <= content_h:
            return float(pt)
    return float(_MIN_ENGLISH_PT)


def _pack_english_into_slides(sections, labels, n_cols, max_rows,
                                text_width_in: float, content_h: float):
    """Pack English sections into slides/columns."""
    slides = []
    i = 0
    n = len(sections)
    # trailing_asts for English currently unused for visual sizing here; the
    # injected chorus copies have n_asts=0 so the asterisk text only adds
    # length on verse last lines, which is rarely the wrap-deciding factor.
    # Use zeros for packing fit; rendering uses the real asts.
    asts_zero = [0] * n
    while i < n:
        slide_cols = _pack_english_slide(sections, labels, asts_zero,
                                          i, n, n_cols, max_rows,
                                          text_width_in, content_h)
        if not slide_cols:
            break
        slides.append(slide_cols)
        i += sum(len(col) for col in slide_cols)
    return slides


def _pack_english_slide(sections, labels, trailing_asts,
                         start, n, n_cols, max_rows,
                         text_width_in: float, content_h: float):
    """Pack one slide worth of English sections into columns. Width-aware:
    uses visual (wrapped) line count to decide what fits."""
    if start >= n:
        return []

    def max_fit(idx):
        count = 0
        while idx + count < n:
            if max_rows and count >= max_rows:
                break
            cand = list(range(idx, idx + count + 1))
            # Find the largest pt within [_MIN_ENGLISH_PT, _MAX_ENGLISH_PT]
            # that fits this candidate column. If even MIN doesn't fit, stop.
            pt = _compute_english_col_font_pt(cand, sections, labels,
                                               trailing_asts, text_width_in,
                                               content_h)
            line_h_in = pt * _ENGLISH_LINE_H / 72.0
            vis = _english_col_visual_lines(cand, sections, labels,
                                             trailing_asts, pt, text_width_in)
            if pt <= _MIN_ENGLISH_PT and vis * line_h_in > content_h:
                break
            count += 1
        return max(count, 1)

    if n_cols == 1:
        k = max_fit(start)
        return [list(range(start, start + k))]

    max_k = max_fit(start)
    best_total = -1
    best_k = max_k
    best_diff = None

    for k in range(1, max_k + 1):
        rest = _pack_english_slide(sections, labels, trailing_asts,
                                    start + k, n, n_cols - 1, max_rows,
                                    text_width_in, content_h)
        rest_total = sum(len(c) for c in rest)
        total = k + rest_total
        right_count = len(rest[0]) if rest else 0
        diff = abs(k - right_count)
        better = (
            total > best_total
            or (total == best_total and (best_diff is None or diff < best_diff))
            or (total == best_total and diff == best_diff and k > best_k)
        )
        if better:
            best_total = total
            best_k = k
            best_diff = diff

    col0 = list(range(start, start + best_k))
    rest = _pack_english_slide(sections, labels, trailing_asts,
                                start + best_k, n, n_cols - 1, max_rows,
                                text_width_in, content_h)
    return [col0] + rest


def _render_english_column(slide, sections, section_indices,
                            labels, trailing_asts,
                            col_left: float, num_margin: float,
                            text_width: float, font_pt: float, config,
                            content_top: float = None,
                            content_h: float = None):
    """Render one English column as a stack of per-section text boxes.

    Each section gets its own text box so verse numbers (placed in a narrow
    left column at the same Y) align precisely with the section's first line.
    Box heights are distributed proportionally across the available column
    height so estimation errors in the visual-line counter don't create
    visible gaps. Between sections we leave a one-line gap. `normAutofit` is
    enabled as a safety net for the rare case where the wrap estimate
    under-counts.
    """
    from pptx.enum.text import MSO_ANCHOR

    if content_top is None:
        content_top = _CONTENT_TOP
    if content_h is None:
        content_h = _CONTENT_H

    line_h_in = font_pt * _ENGLISH_LINE_H / 72.0
    num_font_pt = max(int(round(font_pt * 0.9)), 14)
    rgb = _parse_color(config.text_color)
    line_spacing = Pt(font_pt * _ENGLISH_LINE_H)

    # Proportional height distribution: each section's box gets a share of
    # (content_h minus inter-section gaps) proportional to its visual lines.
    # This ensures sections fill the column evenly regardless of whether the
    # line-wrap estimator over- or under-counts on any individual line.
    n_secs = len(section_indices)
    vis_per_section = [
        _english_section_visual_lines(sections[vi], labels[vi],
                                       trailing_asts[vi], font_pt, text_width)
        for vi in section_indices
    ]
    total_vis = sum(vis_per_section) or 1
    gap_total = max(0, n_secs - 1) * line_h_in
    available_text_h = max(content_h - gap_total, line_h_in * n_secs)

    y = content_top
    for pos, vi in enumerate(section_indices):
        section = sections[vi]
        label   = labels[vi]
        n_asts  = trailing_asts[vi]
        box_h = (vis_per_section[pos] / total_vis) * available_text_h

        text_left = col_left + num_margin
        tx = slide.shapes.add_textbox(
            Inches(text_left), Inches(y),
            Inches(text_width), Inches(box_h))
        tf = tx.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.margin_left = 0
        tf.margin_right = Inches(0.05)
        # Safety net: if the wrap estimate under-counts (e.g. very long word),
        # PowerPoint will shrink this box's text to fit rather than overflow.
        _enable_text_shrink_to_fit(tf)

        first_para = True
        def _add_para(text, bold=False):
            nonlocal first_para
            p = tf.paragraphs[0] if first_para else tf.add_paragraph()
            first_para = False
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            p.line_spacing = line_spacing
            run = p.add_run()
            run.text = text
            run.font.size = Pt(font_pt)
            run.font.bold = bold
            run.font.color.rgb = rgb
            run.font.name = "Arial"

        if label:
            _add_para(label)

        for line_idx, line in enumerate(section.lines):
            line_text = line.words[0].tokens[0].char if line.words else ""
            if line_idx == len(section.lines) - 1 and n_asts > 0:
                line_text += ' ' + ' '.join(['*'] * n_asts)
            _add_para(line_text)

        # First-content-line offset within this box
        first_content_offset = (1 if label else 0) * line_h_in
        if section.number and num_margin > 0:
            num_y = y + first_content_offset
            num_tx = slide.shapes.add_textbox(
                Inches(col_left), Inches(num_y),
                Inches(num_margin), Inches(line_h_in + 0.05))
            num_tf = num_tx.text_frame
            num_tf.word_wrap = False
            num_tf.vertical_anchor = MSO_ANCHOR.TOP
            num_tf.margin_top = 0
            num_tf.margin_bottom = 0
            num_tf.margin_left = 0
            num_tf.margin_right = 0
            p = num_tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            p.line_spacing = line_spacing
            run = p.add_run()
            run.text = f"{section.number}."
            run.font.size = Pt(num_font_pt)
            run.font.color.rgb = rgb
            run.font.name = "Arial"

        y += box_h + line_h_in  # inter-section gap = one blank line


# ── Chinese packing ───────────────────────────────────────────────────────────

def _pack_into_columns(section_imgs, config: SlideConfig,
                       img_w: float, ch_px: int,
                       content_h: float = None,
                       min_char_h: float = None) -> list:
    """Pack sections into columns/slides, returning a list of slides.

    `content_h` overrides the default content area height; pass the per-slide
    available height (which can be smaller when the title wraps).
    `min_char_h` overrides the global _MIN_CHAR_H_IN floor (inches).
    """
    if content_h is None:
        content_h = _CONTENT_H
    if min_char_h is None:
        min_char_h = _MIN_CHAR_H_IN
    max_rows = config.rows_per_column
    slides   = []
    i        = 0
    n        = len(section_imgs)

    while i < n:
        slide_cols = _best_pack_slide(section_imgs, config, img_w, ch_px,
                                      i, n, config.columns, max_rows,
                                      content_h, min_char_h)
        if not slide_cols:
            break
        slides.append(slide_cols)
        i += sum(len(col) for col in slide_cols)

    return slides


def _max_fit_in_col(section_imgs, config: SlideConfig, img_w: float, ch_px: int,
                    start: int, n: int, max_rows,
                    content_h: float = None,
                    min_char_h: float = None) -> int:
    """Return the max number of sections (starting at `start`) that fit in one column."""
    if content_h is None:
        content_h = _CONTENT_H
    if min_char_h is None:
        min_char_h = _MIN_CHAR_H_IN
    col = []
    idx = start
    while idx < n:
        if max_rows and len(col) >= max_rows:
            break
        candidate    = col + [idx]
        rows_in_col  = len(candidate)
        row_h = (content_h - _ROW_GAP * max(0, rows_in_col - 1)) / rows_in_col
        all_fit = all(
            ch_px * min(
                img_w / section_imgs[j].width,
                row_h / section_imgs[j].height,
            ) >= min_char_h
            for j in candidate
        )
        if all_fit:
            col.append(idx)
            idx += 1
        else:
            break
    return len(col)


def _best_pack_slide(section_imgs, config: SlideConfig, img_w: float, ch_px: int,
                     start: int, n: int, n_cols: int, max_rows,
                     content_h: float = None,
                     min_char_h: float = None) -> list:
    """Pack sections into n_cols columns starting at `start`."""
    if content_h is None:
        content_h = _CONTENT_H
    if min_char_h is None:
        min_char_h = _MIN_CHAR_H_IN
    if start >= n:
        return []

    max_col0 = _max_fit_in_col(section_imgs, config, img_w, ch_px,
                                start, n, max_rows, content_h, min_char_h)
    if max_col0 == 0:
        max_col0 = 1

    if n_cols == 1:
        count = min(max_col0, n - start)
        return [list(range(start, start + count))]

    best_total = -1
    best_k     = max_col0
    best_diff  = None

    for k in range(1, max_col0 + 1):
        rest       = _best_pack_slide(section_imgs, config, img_w, ch_px,
                                      start + k, n, n_cols - 1, max_rows,
                                      content_h, min_char_h)
        rest_total = sum(len(col) for col in rest)
        total      = k + rest_total
        # Prefer larger total; for ties, prefer balanced split (smaller
        # |col0 - col1|); for ties in that too, prefer larger k (fill left
        # column first).
        right_count = rest[0] if rest else []
        diff = abs(k - (len(right_count) if isinstance(right_count, list) else 0))
        better = (
            total > best_total
            or (total == best_total and (best_diff is None or diff < best_diff))
            or (total == best_total and diff == best_diff and k > best_k)
        )
        if better:
            best_total = total
            best_k     = k
            best_diff  = diff

    col0 = list(range(start, start + best_k))
    rest = _best_pack_slide(section_imgs, config, img_w, ch_px,
                            start + best_k, n, n_cols - 1, max_rows,
                            content_h, min_char_h)
    return [col0] + rest


def _uniform_scale_for_slide(slide_cols: list, section_imgs,
                              img_w: float, content_h: float = None) -> float:
    """Return the uniform scale that makes all sections on a slide the same size."""
    if content_h is None:
        content_h = _CONTENT_H
    scales = []
    for col_group in slide_cols:
        rows_in_col = len(col_group)
        row_h = (content_h - _ROW_GAP * max(0, rows_in_col - 1)) / rows_in_col
        for vi in col_group:
            img = section_imgs[vi]
            scales.append(min(img_w / img.width, row_h / img.height))
    return min(scales) if scales else 1.0


# ── pptx helpers ─────────────────────────────────────────────────────────────

def _estimate_text_width_in(text: str, pt: int) -> float:
    """Rough rendered width (inches) of `text` at the given point size.

    Chinese chars ~1.0em, ASCII chars ~0.55em (Arial-ish).
    """
    if not text:
        return 0.0
    w = 0.0
    for c in text:
        if ord(c) > 0x3000:
            w += pt * 1.0 / 72.0
        else:
            w += pt * 0.55 / 72.0
    return w


def _estimate_title_width_in(title_zh: str, title_py: str,
                              zh_pt: int, py_pt: int) -> float:
    """Rough width (inches) of the title line at given font sizes.

    Chinese chars ~1.0em, ASCII chars ~0.55em (Arial-ish). Includes the two
    spaces between zh and pinyin runs.
    """
    w_zh = _estimate_text_width_in(title_zh, zh_pt)
    gap = (2 * py_pt * 0.28 / 72.0) if (title_zh and title_py) else 0.0
    w_py = _estimate_text_width_in(title_py, py_pt)
    return w_zh + gap + w_py


_BOOK_PAGE_PT = 14


def _estimate_book_page_width_in(book: str, page: str, pt: int) -> float:
    """Estimated rendered width (inches) of the book/page block.

    The book and page lines stack vertically, right-aligned, so the block's
    width is the max of the two line widths. Returns 0 when both are empty.
    """
    bw = _estimate_text_width_in(book, pt)
    pw = _estimate_text_width_in(page, pt)
    return max(bw, pw)


def _book_page_reservation_in(book: str, page: str) -> float:
    """Inches reserved on the right for the book/page block, including a gap.

    Falls back to the historical 2.0" reservation when book/page is empty so
    the layout doesn't shift unexpectedly for songs without a reference.
    """
    if not book and not page:
        return _TOPRIGHT_RESERVE
    bp_w = _estimate_book_page_width_in(book, page, _BOOK_PAGE_PT)
    # Add a small gap so the title and book/page never visually touch.
    return min(_TOPRIGHT_RESERVE, max(0.5, bp_w + 0.15))


def _add_title_box(slide, title_zh: str, title_py: str, color: str,
                    book: str = '', page: str = '') -> float:
    """Add a left-aligned title text box with Chinese and pinyin runs.

    Shrinks the title font (down to 30 pt) if the default 36 pt would overflow
    the available width. Below that, the title wraps to two lines and the
    text box grows accordingly.

    Returns the actual title-box height used (inches), so callers can push the
    content area down when the title wraps.
    """
    from pptx.enum.text import MSO_ANCHOR
    bp_reserve = _book_page_reservation_in(book, page)
    title_w = max(1.0, _CONTENT_W - bp_reserve)

    zh_pt, py_pt = _TITLE_ZH_PT, _TITLE_PY_PT
    for candidate in (36, 34, 32, 30):
        if _estimate_title_width_in(title_zh, title_py, candidate, candidate) <= title_w:
            zh_pt = candidate
            py_pt = candidate
            break
    else:
        zh_pt = 30
        py_pt = 30

    # If it still overflows at 30pt, wrap to 2 lines and grow the title box.
    wrap = _estimate_title_width_in(title_zh, title_py, zh_pt, py_pt) > title_w
    title_h = _TITLE_H * 2 if wrap else _TITLE_H

    tx = slide.shapes.add_textbox(
        Inches(_MARGIN_L), Inches(_MARGIN_TOP),
        Inches(title_w), Inches(title_h))
    tf = tx.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if not wrap else MSO_ANCHOR.TOP
    _enable_text_shrink_to_fit(tf)

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    rgb = _parse_color(color)

    if title_zh:
        r = p.add_run()
        r.text = title_zh
        r.font.size  = Pt(zh_pt)
        r.font.bold  = False
        r.font.color.rgb = rgb
        r.font.name  = "Arial"

    if title_py:
        r = p.add_run()
        r.text = ("  " if title_zh else "") + title_py
        r.font.size  = Pt(py_pt)
        r.font.bold  = False
        r.font.color.rgb = rgb
        r.font.name  = "Arial"

    return title_h


def _content_top_for_title(title_zh: str, title_py: str,
                             book: str, page: str) -> float:
    """Compute the per-slide content-area top, accounting for title wrap.

    Mirrors the wrap-detection logic in `_add_title_box` so packing decisions
    made before the title is rendered can budget for a doubled title height.
    """
    if not (title_zh or title_py):
        return _MARGIN_TOP + _TITLE_H + _TITLE_GAP
    bp_reserve = _book_page_reservation_in(book, page)
    title_w = max(1.0, _CONTENT_W - bp_reserve)
    chosen_pt = 30
    for candidate in (36, 34, 32, 30):
        if _estimate_title_width_in(title_zh, title_py, candidate, candidate) <= title_w:
            chosen_pt = candidate
            break
    wrap = _estimate_title_width_in(title_zh, title_py, chosen_pt, chosen_pt) > title_w
    title_h = _TITLE_H * 2 if wrap else _TITLE_H
    return _MARGIN_TOP + title_h + _TITLE_GAP


def _add_book_page_box(slide, book: str, page: str, color: str):
    """Add a right-aligned book/page reference in the top-right reserved area.

    The box width shrinks to the estimated text width so it doesn't reserve
    horizontal space the title could use; it stays anchored to the right edge.
    """
    if not book and not page:
        return
    from pptx.enum.text import MSO_ANCHOR

    bp_w = max(0.4, _estimate_book_page_width_in(book, page, _BOOK_PAGE_PT) + 0.05)
    bp_left = _MARGIN_L + _CONTENT_W - bp_w

    tx = slide.shapes.add_textbox(
        Inches(bp_left), Inches(_MARGIN_TOP),
        Inches(bp_w), Inches(_TITLE_H))
    tf = tx.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    rgb = _parse_color(color)
    lines = [l for l in (book, page) if l]
    for i, text in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = text
        r.font.size = Pt(_BOOK_PAGE_PT)
        r.font.bold = False
        r.font.color.rgb = rgb
        r.font.name = "Arial"


def _add_text_box(slide, text: str,
                  left: float, top: float, width: float, height: float,
                  font_pt: int, align, color: str,
                  bold: bool = False, v_anchor: str = 'middle'):
    from pptx.enum.text import MSO_ANCHOR
    tx = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tx.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP if v_anchor == 'top' else MSO_ANCHOR.MIDDLE

    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text

    f = run.font
    f.size  = Pt(font_pt)
    f.bold  = bold
    f.color.rgb = _parse_color(color)
    f.name  = "Arial"


def _place_verse_image(slide, img, left: float, top: float, scale: float):
    """Place a PIL RGBA image using a pre-computed uniform scale (inches/pixel)."""
    buf = BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    slide.shapes.add_picture(
        buf,
        Inches(left), Inches(top),
        Inches(img.width * scale), Inches(img.height * scale),
    )


def _parse_color(hex_color: str) -> RGBColor:
    h = hex_color.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _enable_text_shrink_to_fit(text_frame):
    """Set PowerPoint's 'Shrink text on overflow' (normAutofit) on a text frame.

    python-pptx doesn't expose this directly; we insert the element via lxml.
    With no attributes, PowerPoint computes the shrink scale at display time.
    """
    from pptx.oxml.ns import qn
    from lxml import etree

    bodyPr = text_frame._txBody.find(qn('a:bodyPr'))
    if bodyPr is None:
        return
    for tag in ('a:normAutofit', 'a:spAutoFit', 'a:noAutofit'):
        existing = bodyPr.find(qn(tag))
        if existing is not None:
            bodyPr.remove(existing)
    etree.SubElement(bodyPr, qn('a:normAutofit'))
